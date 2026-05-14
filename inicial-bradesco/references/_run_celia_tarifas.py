"""Inicial TARIFAS — CLIENTE EXEMPLO RODRIGUES DA SILVA. Usa rotina automática que:
1. Detecta extrato digital com text-layer (no `0. Kit/`)
2. Parseia posicionalmente TODOS os lançamentos com 'TARIFA' na descrição
   (CESTA cheia + VR.PARCIAL + EMISSÃO EXTRATO)
3. Compara com a tabela XLSX do NotebookLM e detecta divergência
4. Gera planilha XLSX v2 substituta (formato igual ao NotebookLM, mas completa)
5. Usa esses dados para preencher a inicial

Caso paradigma — quando o NotebookLM ignora rubricas relevantes (VR.PARCIAL,
EMISSÃO EXTRATO), a skill detecta e completa automaticamente.
"""
import io, sys, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from _pipeline_caso import montar_dados_padrao
from helpers_docx import aplicar_template
from extenso import extenso_cardinal
from auditor_tarifas_completo import auditar_e_completar_tarifas, lancamentos_para_tese

PASTA_CLIENTE = r'C:\Users\gabri\OneDrive\Área de Trabalho\APP - BRADESCO\0. TESTE 1\CLIENTE EXEMPLO RODRIGUES DA SILVA - Maurivã - TARIFAS'
PASTA_ACAO = os.path.join(PASTA_CLIENTE, 'TARIFA')
TABELA_NB = os.path.join(PASTA_CLIENTE, 'Tabela de Descontos por Procuracao - CLIENTE EXEMPLO RODRIGUES DA SILVA.xlsx')
TEMPLATE = r'C:\Users\gabri\OneDrive\Documentos\Obsidian Vault\Modelos\IniciaisBradesco\_templates\inicial-tarifas.docx'
DOCX_OUT = os.path.join(PASTA_ACAO, 'INICIAL_Tarifas_CLIENTE EXEMPLO_v3.docx')

# 1. Auditoria + completar TUDO de TARIFA
audit = auditar_e_completar_tarifas(
    pasta_cliente=PASTA_CLIENTE,
    tabela_xlsx_path=TABELA_NB,
    cliente_nome='CLIENTE EXEMPLO RODRIGUES DA SILVA',
    conta_label='Agência: 3706 | Conta: 16649-9',
    procuracao_label='TARIFA BANCÁRIA - CESTA B.EXPRESSO',
    gerar_planilha_v2=True,
)
print(f'Severidade: {audit["severidade"]}')
print(f'Extrato direto: {audit["qtd_extrato"]} lanç / R$ {audit["soma_extrato"]:.2f}')
print(f'XLSX original:  {audit["qtd_xlsx_original"]} lanç')
if audit['planilha_v2_path']:
    print(f'Planilha v2: {os.path.basename(audit["planilha_v2_path"])}')

LANCAMENTOS = lancamentos_para_tese(audit['lancamentos'])

# 2. Qualificação extraída da notificação extrajudicial
autora = {
    'nome': 'CLIENTE EXEMPLO RODRIGUES DA SILVA', 'nacionalidade': 'brasileira',
    'estado_civil': '', 'profissao': 'aposentada',
    'cpf': '000.000.003-13', 'rg': '1000002-2',
    'orgao_expedidor_prefixo': 'SSP/AM',
    'logradouro': 'Rua 10', 'numero': '818', 'bairro': 'Nova Esperança',
    'cidade': 'Maués', 'cep': '69.190-000',
}
conta = {'agencia': '3706', 'numero': '16649-9'}
renda = {'valor_float': 980.09}

tese = {
    'rubrica': 'TARIFA BANCÁRIA - CESTA B.EXPRESSO',
    'lancamentos': LANCAMENTOS,
}
terceiro = {'nome': '', 'cnpj': '', 'logradouro': '', 'numero': '',
            'bairro': '', 'cidade': '', 'uf': '', 'cep': ''}

dados, calc = montar_dados_padrao(
    autora=autora, conta=conta, renda=renda, tese=tese, terceiro=terceiro,
    eh_idoso=False, competência='Maués', uf='AM',
)
dados['titulo'] = ('TARIFA BANCÁRIA - CESTA B.EXPRESSO / TARIFA BANCÁRIA - VR.PARCIAL CESTA / '
                   'TARIFA EMISSÃO EXTRATO')
dados['remuneração'] = 'aposentadoria pelo INSS'
dados['numero_desconto'] = str(len(LANCAMENTOS))
dados['desconto_extenso'] = extenso_cardinal(len(LANCAMENTOS))

print(f'\nGerando inicial: {len(LANCAMENTOS)} lanç / R$ {calc["total"]:.2f} / VC R$ {calc["valor_causa"]:.2f}')
res = aplicar_template(TEMPLATE, dados, DOCX_OUT)
print(f'Template: {res["modificados"]} modif, residuais: {res["residuais"] or "nenhum"}')

# 3. Pós-fix: placeholders extras + bug R$ R$
import zipfile, re
with zipfile.ZipFile(DOCX_OUT, 'r') as z:
    nomes = z.namelist()
    buf = {n: z.read(n) for n in nomes}
xml = buf['word/document.xml'].decode('utf-8')
for k, v in [('{{remuneração}}', dados['remuneração']),
             ('{{valor_remuneração}}', dados['valor_remuneração']),
             ('{{valor_remuneração_extenso}}', dados['valor_remuneração_extenso'])]:
    xml = xml.replace(k, v)
buf['word/document.xml'] = xml.encode('utf-8')
os.remove(DOCX_OUT)
with zipfile.ZipFile(DOCX_OUT, 'w', zipfile.ZIP_DEFLATED) as z:
    for n in nomes: z.writestr(n, buf[n])

# Pós-fix R$ R$ via python-docx (junta runs)
from docx import Document
d = Document(DOCX_OUT)
n_fix = 0
for p in d.paragraphs:
    if 'R$' in p.text and re.search(r'R\$\s*R\$', p.text):
        full = ''.join(r.text for r in p.runs)
        novo = re.sub(r'R\$\s*R\$', 'R$', full)
        if full != novo:
            p.runs[0].text = novo
            for r in p.runs[1:]: r.text = ''
            n_fix += 1
d.save(DOCX_OUT)
print(f'Pós-fix R$ R$: {n_fix} parágrafos')
print(f'OK -> {DOCX_OUT}')
