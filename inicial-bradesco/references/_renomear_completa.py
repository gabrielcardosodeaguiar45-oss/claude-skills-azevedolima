# -*- coding: utf-8 -*-
"""Renomeia placeholders antigos {{rubrica_completa(_caps)}} para o novo padrão
{{rubrica_mora_canonica(_caps)}} em inicial-mora.docx."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent))
from docx import Document
from helpers_docx import substituir_in_run

caminho = Path(r"C:\Users\gabri\OneDrive\Documentos\Obsidian Vault\Modelos\IniciaisBradesco\_templates\inicial-mora.docx")

MAPA = {
    "{{rubrica_completa_caps}}": "{{rubrica_mora_canonica_caps}}",
    "{{rubrica_completa}}":      "{{rubrica_mora_canonica}}",
}

doc = Document(caminho)
n = 0
for p in doc.paragraphs:
    if substituir_in_run(p._p, MAPA):
        n += 1
doc.save(caminho)
print(f"Parágrafos modificados: {n}")
