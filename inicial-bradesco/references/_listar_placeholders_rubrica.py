# -*- coding: utf-8 -*-
"""Lista todos os placeholders {{rubrica_*}} em cada template e quantas vezes
aparecem, para revisão."""
import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pathlib import Path
from docx import Document
from collections import Counter

VAULT = Path(r"C:\Users\gabri\OneDrive\Documentos\Obsidian Vault\Modelos\IniciaisBradesco\_templates")

PAT = re.compile(r"\{\{rubrica_[a-z_]+\}\}")

for nome in ["inicial-mora.docx", "inicial-mora-encargo.docx", "inicial-combinada.docx"]:
    caminho = VAULT / nome
    if not caminho.exists():
        continue
    print(f"\n=== {nome} ===")
    doc = Document(caminho)
    contador = Counter()
    for p in doc.paragraphs:
        for m in PAT.finditer(p.text):
            contador[m.group(0)] += 1
    if not contador:
        print("  (nenhum placeholder de rubrica)")
        continue
    for ph, n in sorted(contador.items()):
        print(f"  {ph} × {n}")
