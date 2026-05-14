"""Validador automático de template novo.

Verifica se um template DOCX tem os elementos estruturais que a skill espera
(cabeçalho, qualificação, polo passivo, síntese fática, bloco fático, pedidos,
rodapé, valor da causa).

USO:
    python validar_template.py inicial-jfpe-base.docx
    python validar_template.py /caminho/completo/template.docx

SAÍDA:
    ✓ ou ⚠ ou 🚨 para cada item da checklist.

Ajuda a NÃO ESQUECER NADA ao adicionar uma nova UF.
"""
import sys
import os
import re
from typing import List, Dict
from docx import Document


W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

# Cada item: (tag, descrição, padrão_regex_OU_substring, severidade)
# severidade: 'critico' | 'recomendado' | 'opcional'
CHECKLIST = [
    # Cabeçalho (tolera "Ao Juízo" ou só "Juízo" — alguns templates omitem o "Ao")
    ('cabecalho_juizo',
     'Cabeçalho "Ao Juízo..." (ou "Juízo...")',
     r'(Ao\s+)?Ju[íi]zo (do|da|de)',
     'critico'),

    # Prioridade idoso
    ('prioridade_cabecalho',
     'Cabeçalho prioridade idoso (CPC 1.048)',
     r'Prioridade de tramita[çc][ãa]o.*1\.?048',
     'recomendado'),

    # Ementa
    ('ementa',
     'EMENTA DO CASO',
     r'EMENTA DO CASO',
     'recomendado'),

    # Qualificação do autor — placeholder OU nome em CAPS antes de
    # 'brasileiro/brasileira' (padrão de qualificação do escritório)
    ('qualif_autor',
     'Parágrafo de qualificação do autor (placeholder ou nome piloto)',
     r'(\{\{nome_(autor|completo)\}\}|[A-ZÀ-Ú]{3,}(?:\s[A-ZÀ-Ú]{2,})+\s*,\s*brasileir[oa])',
     'critico'),

    # Pedido principal
    ('pedido_acao',
     'AÇÃO DECLARATÓRIA DE INEXISTÊNCIA',
     r'A[ÇC][ÃA]O DECLARAT[ÓO]RIA DE INEXIST[ÊE]NCIA',
     'critico'),

    # Polo passivo
    ('polo_passivo',
     'Polo passivo "em face de"',
     r'em face de',
     'critico'),

    # Síntese fática
    ('sintese_fatica',
     'Síntese fática "recebe benefício previdenciário"',
     r'recebe benef[íi]cio previdenci[áa]rio',
     'critico'),

    # Intro contratos (várias variações possíveis nos templates)
    ('intro_contratos',
     'Intro "tomou conhecimento" ou "constatou a existência" dos descontos',
     r'(tomou conhecimento dos descontos|constatou a exist[êe]ncia de descontos|ao verificar.*descontos)',
     'critico'),

    # Bloco fático
    ('bloco_fatico',
     'Bloco fático "No que diz respeito"',
     r'No que diz respeito (ao|aos) referid',
     'critico'),

    # DO DIREITO
    ('do_direito',
     'Seção DO DIREITO',
     r'DO DIREITO',
     'recomendado'),

    # CDC
    ('cdc',
     'Aplicação do CDC',
     r'C[óo]digo de Defesa do Consumidor|CDC',
     'recomendado'),

    # Pedidos
    ('pedidos_titulo',
     'DOS PEDIDOS',
     r'DOS PEDIDOS',
     'critico'),

    ('pedido_declaratorio',
     'Pedido "Declarar a inexistência do empréstimo/refinanciamento" (singular ou plural)',
     r'(D|d)eclarar a inexist[êe]ncia (do|dos)\s+(seguintes\s+)?(empr[ée]stimo|refinanciamento)',
     'critico'),

    ('pedido_dano_moral',
     'Pedido de danos morais (R$ 15.000 ou similar)',
     r'(danos morais|indeniza[çc][ãa]o por danos)',
     'critico'),

    # Justiça gratuita
    ('justica_gratuita',
     'Pedido Justiça Gratuita',
     r'Justi[çc]a Gratuita',
     'recomendado'),

    # Rodapé
    ('valor_causa',
     'Valor da causa "Dá-se a causa"',
     r'D[áa]-se a causa',
     'critico'),

    ('cidade_data',
     'Cidade/UF + data (ou placeholder)',
     r'(/(AL|AM|BA|SE|MG|ES|SC|PE|RJ|SP)\s*,\s*(\d|de)|\{\{(cidade|comarca|data)[^}]*\}\})',
     'recomendado'),

    ('procurador_assinatura',
     'Assinatura do procurador (Tiago / Patrick / Gabriel / Alexandre / Eduardo)',
     r'(Tiago de Azevedo|Patrick Willian|Gabriel Cardoso|Alexandre|Eduardo Fernando)',
     'critico'),

    ('procurador_oab',
     'OAB do procurador',
     r'OAB/(AL|AM|BA|SE|MG|ES|SC|PE)\s*\d+[A-Z]?',
     'critico'),
]


def validar(path_template: str) -> Dict:
    """Valida o template e retorna dict com achados.

    Returns:
        {
            'arquivo': str,
            'total_paragrafos': int,
            'achados': [
                {'tag': ..., 'desc': ..., 'severidade': ...,
                 'encontrado': bool, 'paragrafo': int (ou None)}
            ],
            'criticos_faltando': int,
            'recomendados_faltando': int,
        }
    """
    if not os.path.exists(path_template):
        return {'erro': f'Arquivo não existe: {path_template}'}

    doc = Document(path_template)
    pars = list(doc.paragraphs)

    achados = []
    criticos_faltando = 0
    recomendados_faltando = 0

    for tag, desc, padrao, severidade in CHECKLIST:
        encontrado = False
        idx = None
        for ip, p in enumerate(pars):
            if re.search(padrao, p.text):
                encontrado = True
                idx = ip
                break

        achados.append({
            'tag': tag,
            'desc': desc,
            'severidade': severidade,
            'encontrado': encontrado,
            'paragrafo': idx,
        })

        if not encontrado:
            if severidade == 'critico':
                criticos_faltando += 1
            elif severidade == 'recomendado':
                recomendados_faltando += 1

    # Detecções extras informativas
    extras = {
        'tem_inss_polo_passivo':
            any('INSTITUTO NACIONAL DO SEGURO SOCIAL' in p.text for p in pars),
        'tem_2_bancos_polo_passivo':
            any(p.text.count('em face de') == 1 and 'e BANCO' in p.text and 'INSTITUTO' in p.text
                for p in pars),
        'tem_marcador_se_idoso':
            any('{{SE_IDOSO}}' in p.text for p in pars),
        'tem_blocos_pre_prontos_AL':
            sum(1 for p in pars if any(x in p.text for x in
                ['(1 CONTRATO,', '(2 CONTRATOS,', '(1 REFIN,', '(2 REFIN,'])),
        'placeholders_unicos': sorted({m for p in pars
                                         for m in re.findall(r'\{\{[^}]+\}\}', p.text)}),
    }

    return {
        'arquivo': path_template,
        'total_paragrafos': len(pars),
        'achados': achados,
        'criticos_faltando': criticos_faltando,
        'recomendados_faltando': recomendados_faltando,
        'extras': extras,
    }


def imprimir_relatorio(resultado: Dict):
    """Formata e imprime o relatório de validação."""
    if 'erro' in resultado:
        print(f'❌ {resultado["erro"]}')
        return

    print(f'=== VALIDAÇÃO DE TEMPLATE ===')
    print(f'Arquivo: {os.path.basename(resultado["arquivo"])}')
    print(f'Total parágrafos: {resultado["total_paragrafos"]}')
    print()
    print('CHECKLIST:')
    for a in resultado['achados']:
        icone_sev = {'critico': '🚨', 'recomendado': '⚠', 'opcional': 'ℹ'}[a['severidade']]
        marca = '✓' if a['encontrado'] else f'{icone_sev} FALTA'
        idx_str = f'(p.{a["paragrafo"]})' if a['encontrado'] and a['paragrafo'] is not None else ''
        print(f'  {marca:12} {a["desc"]:60} {idx_str}')

    print()
    print('EXTRAS:')
    for k, v in resultado['extras'].items():
        if k == 'placeholders_unicos':
            print(f'  {k}: {len(v)} placeholders ' + (f'({v[:6]}...)' if len(v) > 6 else f'({v})'))
        else:
            print(f'  {k}: {v}')

    print()
    print('RESUMO:')
    print(f'  🚨 Críticos faltando: {resultado["criticos_faltando"]}')
    print(f'  ⚠  Recomendados faltando: {resultado["recomendados_faltando"]}')
    if resultado['criticos_faltando'] == 0 and resultado['recomendados_faltando'] == 0:
        print('  ✅ TEMPLATE COMPLETO — pronto para uso')
    elif resultado['criticos_faltando'] == 0:
        print('  ⚠ Template com pendências MENORES — usável mas conferir')
    else:
        print('  🚨 Template INCOMPLETO — corrigir antes de usar')


if __name__ == '__main__':
    import io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

    if len(sys.argv) < 2:
        # Validar TODOS os templates do vault como sanity check
        VAULT = r'C:\Users\gabri\OneDrive\Documentos\Obsidian Vault\Modelos\IniciaisNaoContratado\_templates'
        templates = sorted(f for f in os.listdir(VAULT)
                            if f.endswith('.docx') and not f.endswith('.backup_12blocos'))
        for t in templates:
            print()
            print('=' * 70)
            imprimir_relatorio(validar(os.path.join(VAULT, t)))
    else:
        path = sys.argv[1]
        # Se for nome relativo, procurar no vault
        if not os.path.isabs(path):
            VAULT = r'C:\Users\gabri\OneDrive\Documentos\Obsidian Vault\Modelos\IniciaisNaoContratado\_templates'
            cand = os.path.join(VAULT, path)
            if os.path.exists(cand):
                path = cand
        imprimir_relatorio(validar(path))
