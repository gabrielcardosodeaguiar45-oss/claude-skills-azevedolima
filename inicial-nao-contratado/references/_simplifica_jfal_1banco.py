"""Edita o template inicial-jfal-1banco.docx removendo os 12 blocos fáticos
pré-prontos e substituindo por 1 bloco genérico (decisão Gabriel 07/05/2026,
Opção B). O bloco genérico:
  - Não menciona depósito (neutro — regra do escritório)
  - Tem placeholders xxxxxxxx, xxx,xx, contrato n° xxxxxxx
  - É preenchido dinamicamente pelo pipeline AL conforme N de contratos
"""
import sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from copy import deepcopy

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
TEMPLATE = r'C:\Users\gabri\OneDrive\Documentos\Obsidian Vault\Modelos\IniciaisNaoContratado\_templates\inicial-jfal-1banco.docx'
BACKUP = TEMPLATE + '.backup_12blocos'

# Conteúdo do bloco fático genérico (5 parágrafos, neutro, sem depósito)
NOVO_BLOCO = [
    'No que diz respeito ao referido empréstimo, cumpre informar que a '
    'primeira parcela descontada do benefício da parte autora foi na '
    'competência xxxxxxxx, de um total de xx parcelas, no valor de R$ xxx,xx '
    '(valor por extenso), relativas a um empréstimo consignado no valor de '
    'R$ xxx,xx (valor por extenso), contrato n° xxxxxxx, cuja operação foi '
    'realizada pelo banco xxxxx, ora requerido.',
    'Após tomar conhecimento de tal fato, entrou em contato com o banco '
    'através da Central de Atendimento, contudo, não obteve êxito. Importa '
    'destacar que a parte autora não possui conhecimentos tecnológicos '
    'suficientes para realizar solicitações na modalidade on-line.',
    'Ocorre, Excelência, que a parte autora não contratou tal empréstimo, '
    'e não foi autorizada nenhuma forma de empréstimo consignado, nem mesmo '
    'na modalidade de crédito em conta, mas desde então vem sofrendo com os '
    'descontos indevidos.',
    'Destarte, Excelência, é muito cômodo o banco requerido realizar '
    'empréstimos para pensionistas e aposentados, onde o risco é baixíssimo '
    'e os lucros são exorbitantes.',
    'A parte autora não aceita ter que pagar por um contrato que foi pactuado '
    'sem o seu consentimento.',
]


def main():
    if not os.path.exists(BACKUP):
        import shutil
        shutil.copy(TEMPLATE, BACKUP)
        print(f'Backup criado em: {BACKUP}')
    else:
        print(f'Backup já existe (não sobrescrevendo): {BACKUP}')

    doc = Document(TEMPLATE)
    pars = list(doc.paragraphs)

    # 1. Identificar idx_inicio (primeira âncora) e idx_fim (parágrafo "Sabe-se")
    ANCORAS = [
        '(1 CONTRATO,', '(2 CONTRATOS,', '(1 REFIN,', '(2 REFIN,',
    ]
    SENTINELA = 'Sabe-se que tal fato ocorre'
    idx_inicio = None
    idx_fim = None  # exclusivo
    for ip, par in enumerate(pars):
        if idx_inicio is None and any(a in par.text for a in ANCORAS):
            idx_inicio = ip
        if SENTINELA in par.text:
            idx_fim = ip
            break

    if idx_inicio is None or idx_fim is None:
        print(f'❌ Falhou: idx_inicio={idx_inicio}, idx_fim={idx_fim}')
        return

    print(f'Removendo parágrafos {idx_inicio} a {idx_fim - 1} '
          f'({idx_fim - idx_inicio} parágrafos no total)')

    # 2. Pegar o pPr do PRIMEIRO bloco antigo para usar como modelo
    primeiro_par = pars[idx_inicio]._element
    pPr_modelo = primeiro_par.find(W + 'pPr')
    pPr_modelo_xml = deepcopy(pPr_modelo) if pPr_modelo is not None else None

    # 3. Remover os parágrafos [idx_inicio, idx_fim)
    body = primeiro_par.getparent()
    # Capturar referência do "Sabe-se" para inserir ANTES dele
    sentinela_par = pars[idx_fim]._element

    for ip in range(idx_fim - 1, idx_inicio - 1, -1):
        body.remove(pars[ip]._element)

    # 4. Inserir os 5 parágrafos novos ANTES da sentinela.
    # IMPORTANTE: addprevious mantém a ORDEM de inserção quando o ponto de
    # referência é fixo (a sentinela). Iterar na ORDEM normal (sem reversed).
    for texto in NOVO_BLOCO:
        novo_par = OxmlElement('w:p')
        if pPr_modelo_xml is not None:
            novo_par.append(deepcopy(pPr_modelo_xml))
        r = OxmlElement('w:r')
        rpr = OxmlElement('w:rPr')
        r.append(rpr)
        rfonts = OxmlElement('w:rFonts')
        rfonts.set(qn('w:ascii'), 'Cambria')
        rfonts.set(qn('w:hAnsi'), 'Cambria')
        rpr.append(rfonts)
        t = OxmlElement('w:t')
        t.text = texto
        t.set(qn('xml:space'), 'preserve')
        r.append(t)
        novo_par.append(r)
        sentinela_par.addprevious(novo_par)

    doc.save(TEMPLATE)
    print(f'✓ Template editado e salvo: {TEMPLATE}')
    print(f'  - 12 blocos pré-prontos REMOVIDOS')
    print(f'  - 1 bloco genérico INSERIDO ({len(NOVO_BLOCO)} parágrafos)')


if __name__ == '__main__':
    main()
